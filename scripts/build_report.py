#!/usr/bin/env python3

import re
import sys
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "INFORME_TECNICO.md"
OUTPUT = ROOT / "output" / "Informe_Tecnico_Banco_XYZ.docx"
DOCS = ROOT / "docs"
DIAGRAMS = [
    DOCS / "diagramas" / "01-arquitectura-general.png",
    DOCS / "diagramas" / "02-flujo-batch.png",
    DOCS / "diagramas" / "03-seguridad-oauth.png",
    DOCS / "diagramas" / "04-transferencia.png",
    DOCS / "diagramas" / "05-retiro-atm.png",
    DOCS / "diagramas" / "06-despliegue-local.png",
    DOCS / "diagramas" / "07-arquitectura-aws.png",
]

NAVY = "17365D"
PALE_BLUE = "EDF3F8"
LIGHT_GRAY = "D9E0E7"
TEXT = RGBColor(25, 34, 48)
MUTED = RGBColor(82, 96, 112)


def clean_text(value: str) -> str:
    return (
        value.replace("\u2014", "-")
        .replace("\u2013", "-")
        .replace("\u2011", "-")
        .replace("\u00a0", " ")
    )


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), LIGHT_GRAY)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    run.font.size = Pt(9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def add_inline_runs(paragraph, text):
    text = clean_text(text)
    pattern = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|<https?://[^>]+>)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(23, 73, 111)
        elif token.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            run = paragraph.add_run(f"{label} ({url})")
            run.font.color.rgb = RGBColor(23, 105, 170)
        else:
            url = token[1:-1]
            run = paragraph.add_run(url)
            run.font.color.rgb = RGBColor(23, 105, 170)
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def add_picture(document, image_path: Path, alt_text: str, max_width=6.45, max_height=8.3):
    with Image.open(image_path) as img:
        width_px, height_px = img.size
    ratio = width_px / height_px
    width = max_width
    height = width / ratio
    if height > max_height:
        height = max_height
        width = height * ratio
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(image_path), width=Inches(width), height=Inches(height))
    inline = document.inline_shapes[-1]._inline
    doc_pr = inline.docPr
    doc_pr.set("name", alt_text[:120])
    doc_pr.set("descr", alt_text[:250])


def style_document(document):
    section = document.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.different_first_page_header_footer = True

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    title = document.styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(12)

    for name, size, before, after in (
        ("Heading 1", 17, 12, 7),
        ("Heading 2", 13.5, 10, 5),
        ("Heading 3", 11.5, 8, 4),
    ):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(10.5)

    header = section.header
    hp = header.paragraphs[0]
    hp.text = "Banco XYZ   |   Informe técnico"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.runs[0].font.name = "Arial"
    hp.runs[0].font.size = Pt(8.5)
    hp.runs[0].font.color.rgb = MUTED
    add_page_number(section.footer.paragraphs[0])


def add_cover(document, metadata):
    for _ in range(4):
        document.add_paragraph()
    # The built-in Word Title style can add a decorative border depending on
    # the local Office theme, so the cover title is formatted explicitly.
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Informe técnico de modernización de la plataforma Banco XYZ")
    title_run.bold = True
    title_run.font.name = "Arial"
    title_run.font.size = Pt(25)
    title_run.font.color.rgb = RGBColor(0, 0, 0)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(30)
    run = subtitle.add_run("Procesamiento batch, microservicios, canales BFF, seguridad distribuida y mensajería asíncrona")
    run.font.name = "Arial"
    run.font.size = Pt(13)
    run.font.color.rgb = MUTED

    table = document.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for key in ("Institución", "Asignatura", "Evaluación", "Estudiante", "Docente", "Sección", "Fecha", "Repositorio"):
        row = table.add_row()
        row.cells[0].width = Cm(4.2)
        row.cells[1].width = Cm(10.6)
        row.cells[0].text = key
        value = clean_text(metadata.get(key, ""))
        if value.startswith("<") and value.endswith(">"):
            value = value[1:-1]
        row.cells[1].text = value
        row.cells[0].paragraphs[0].runs[0].bold = True
        for cell in row.cells:
            set_cell_margins(cell, 90, 120, 90, 120)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
    set_table_borders(table)
    document.add_page_break()


def add_contents(document, headings):
    document.add_heading("Contenido", level=1)
    intro = document.add_paragraph("El informe se organiza en las siguientes secciones:")
    intro.paragraph_format.space_after = Pt(8)
    for heading in headings:
        p = document.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        add_inline_runs(p, heading)
    document.add_page_break()


def add_table(document, rows):
    if not rows:
        return
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline_runs(paragraph, value.strip())
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.size = Pt(9)
                if r_idx == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            if r_idx == 0:
                set_cell_shading(cell, NAVY)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
    set_repeat_table_header(table.rows[0])
    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_code_block(document, text):
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F3F5F7")
    set_cell_margins(cell, 120, 150, 120, 150)
    set_table_borders(table)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(clean_text(text.rstrip()))
    run.font.name = "Courier New"
    run.font.size = Pt(8.2)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def parse_metadata(lines):
    metadata = {}
    for line in lines:
        if line.startswith("|") and "---" not in line:
            cells = [cell.strip().replace("**", "") for cell in line.strip().strip("|").split("|")]
            if len(cells) == 2 and cells[0] != "Antecedente":
                metadata[cells[0]] = cells[1]
    return metadata


def build():
    markdown = SOURCE.read_text(encoding="utf-8")
    lines = markdown.splitlines()
    metadata = parse_metadata(lines[:18])
    start = next(i for i, line in enumerate(lines) if line.strip() == "## Resumen ejecutivo")
    headings = [
        clean_text(line[3:].strip())
        for line in lines[start:]
        if line.startswith("## ") and not line.startswith("### ")
    ]

    document = Document()
    style_document(document)
    add_cover(document, metadata)
    add_contents(document, headings)

    i = start
    diagram_index = 0
    paragraph_buffer = []

    def flush_paragraph():
        nonlocal paragraph_buffer
        if paragraph_buffer:
            p = document.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0.55)
            add_inline_runs(p, " ".join(item.strip() for item in paragraph_buffer))
            paragraph_buffer = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            flush_paragraph()
            i += 1
            continue
        if stripped == "---":
            flush_paragraph()
            i += 1
            continue
        if line.startswith("## "):
            flush_paragraph()
            heading = clean_text(line[3:].strip())
            if heading.startswith("Anexo") or heading in {"6. Arquitectura de la solución", "7. Modernización del procesamiento batch", "15. Preparación y estrategia de despliegue en AWS", "16. Estrategia de pruebas", "21. Conclusiones"}:
                if len(document.paragraphs) > 0:
                    document.add_page_break()
            document.add_heading(heading, level=1)
            i += 1
            continue
        if line.startswith("### "):
            flush_paragraph()
            document.add_heading(clean_text(line[4:].strip()), level=2)
            i += 1
            continue
        if stripped.startswith("```"):
            flush_paragraph()
            language = stripped[3:].strip()
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            if language == "mermaid":
                add_picture(document, DIAGRAMS[diagram_index], f"Diagrama técnico {diagram_index + 1}")
                diagram_index += 1
            else:
                add_code_block(document, "\n".join(block))
            continue
        if stripped.startswith("!["):
            flush_paragraph()
            match = re.match(r"!\[([^]]*)\]\(([^)]+)\)", stripped)
            if match:
                alt, relative = match.groups()
                add_picture(document, DOCS / relative, clean_text(alt))
            i += 1
            continue
        if stripped.startswith("|"):
            flush_paragraph()
            raw_rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                candidate = lines[i].strip()
                if not re.match(r"^\|?\s*:?-{3,}", candidate):
                    cells = [cell.strip() for cell in candidate.strip("|").split("|")]
                    if not all(re.fullmatch(r":?-+:?", cell) for cell in cells):
                        raw_rows.append(cells)
                i += 1
            add_table(document, raw_rows)
            continue
        if re.match(r"^-\s+", stripped):
            flush_paragraph()
            p = document.add_paragraph(style="List Bullet")
            add_inline_runs(p, re.sub(r"^-\s+", "", stripped))
            i += 1
            continue
        if re.match(r"^\d+\.\s+", stripped):
            flush_paragraph()
            # Preserve the number written in Markdown. Word's built-in list
            # style otherwise continues numbering across unrelated sections.
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.65)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            p.paragraph_format.space_after = Pt(2)
            add_inline_runs(p, stripped)
            i += 1
            continue
        if stripped.startswith("**Figura"):
            flush_paragraph()
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            add_inline_runs(p, stripped)
            for run in p.runs:
                run.font.size = Pt(9)
            i += 1
            continue
        paragraph_buffer.append(stripped)
        i += 1

    flush_paragraph()
    document.core_properties.title = "Informe técnico de modernización de la plataforma Banco XYZ"
    document.core_properties.subject = "Evaluación Final Transversal de Desarrollo Backend III"
    document.core_properties.author = metadata.get("Estudiante", "Cristian Nahuas")
    document.core_properties.keywords = "Spring Batch, microservicios, BFF, OAuth2, Kafka, Docker, AWS"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    try:
        build()
    except Exception as exc:
        print(f"Error al generar el informe: {exc}", file=sys.stderr)
        raise
